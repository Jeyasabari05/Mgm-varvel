"""Editable Excel layer: generate, edit, import, PDF uses confirmed values."""

from __future__ import annotations

import unittest
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from mgm_selector.catalogue import Catalogue, display_value  # noqa: E402
from mgm_selector.excel_report import (  # noqa: E402
    DATA_START_ROW,
    SHEET_CHANGES,
    SHEET_DS,
    SHEET_EDITABLE,
    SHEET_ORIGINAL,
    SHEET_SOURCE,
    DS_FIELD_START_ROW,
    apply_excel_to_report,
    generate_excel,
    load_excel,
    selection_fingerprint,
    validate_rows,
)
from mgm_selector.pdf_report import generate_pdf  # noqa: E402
from mgm_selector.selection import SelectionEngine  # noqa: E402

try:
    import fitz
except ImportError:  # pragma: no cover
    fitz = None


class ExcelWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = SelectionEngine(Catalogue())
        cls.scratch = ROOT / "scratch"
        cls.scratch.mkdir(exist_ok=True)

    def _report(self):
        result = self.engine.lookup("RF", 1.5, 36.53)
        self.assertTrue(result["ok"], result)
        return result["report"]

    def _val(self, report, key):
        for field in report["fields"]:
            if field["key"] == key:
                return field["value"]
        self.fail(key)

    def test_excel_contains_ds_fields_only_plus_identity(self):
        report = self._report()
        path = self.scratch / "test_editable.xlsx"
        generate_excel(report, path)
        wb = load_workbook(path)
        self.assertEqual(
            set(wb.sheetnames),
            {SHEET_DS, SHEET_EDITABLE, SHEET_ORIGINAL, SHEET_CHANGES, SHEET_SOURCE},
        )
        self.assertEqual(wb.active.title, SHEET_DS)
        self.assertTrue(all(wb[name].sheet_state == "hidden" for name in
                            (SHEET_EDITABLE, SHEET_ORIGINAL, SHEET_CHANGES, SHEET_SOURCE)))
        ds = wb[SHEET_DS]
        self.assertEqual(ds["A1"].value, "PRODUCT INFORMATION")
        self.assertEqual(ds["B8"].value, "DS-021-0046")
        labels = []
        values = []
        for offset, field in enumerate(report["fields"]):
            row = DS_FIELD_START_ROW + offset
            labels.append(ds.cell(row, 1).value)
            values.append(ds.cell(row, 2).value)
        expected_labels = [
            "Flange diameter (If flange mounted)" if f["key"] == "flange_diameter" else f["label"]
            for f in report["fields"]
        ]
        expected_values = [
            (": —" if f["key"] == "additional_features" and not f["value"]
             else display_value(f).replace(") : ", "): "))
            for f in report["fields"]
        ]
        self.assertEqual(labels, expected_labels)
        self.assertEqual(values, expected_values)
        self.assertEqual(ds["A4"].value, "CATALOG DESIGNATION")
        self.assertEqual(ds["A5"].value, report["catalog_designation"])
        self.assertEqual(ds["A6"].value, report["product_family"])
        self.assertEqual(len(ds._images), 4)
        torque_row_on_page = DS_FIELD_START_ROW + [f["key"] for f in report["fields"]].index("output_torque")
        self.assertFalse(ds.cell(torque_row_on_page, 2).protection.locked)
        ws = wb[SHEET_EDITABLE]
        keys = []
        row = DATA_START_ROW
        while ws.cell(row, 1).value:
            keys.append(ws.cell(row, 1).value)
            row += 1
        ds_keys = [f["key"] for f in report["fields"]]
        self.assertIn("catalog_designation", keys)
        self.assertEqual(keys[-len(ds_keys) :], ds_keys)
        self.assertNotIn("P1max", keys)
        self.assertNotIn("Gd2", keys)
        torque_row = keys.index("output_torque") + DATA_START_ROW
        self.assertFalse(ws.cell(torque_row, 3).protection.locked)
        self.assertTrue(ws.cell(torque_row, 2).protection.locked)

    def test_edited_excel_is_not_overwritten_and_pdf_uses_it(self):
        report = self._report()
        original_torque = self._val(report, "output_torque")
        self.assertEqual(original_torque, "349")
        path = self.scratch / "test_edited.xlsx"
        generate_excel(report, path)

        wb = load_workbook(path)
        ws = wb[SHEET_DS]
        torque_row = DS_FIELD_START_ROW + [f["key"] for f in report["fields"]].index("output_torque")
        ws.cell(torque_row, 2).value = "(Nm): 337.5"
        wb.save(path)

        loaded = load_excel(path)
        required = [f["key"] for f in report["fields"]]
        check = validate_rows(loaded["editable"], loaded["original"], required)
        self.assertTrue(check["ok"], check["errors"])
        edited_keys = {e["key"] for e in check["edited"]}
        self.assertIn("output_torque", edited_keys)

        confirmed = apply_excel_to_report(report, loaded)
        self.assertEqual(self._val(confirmed, "output_torque"), "337.5")
        # Catalogue report object is unchanged.
        self.assertEqual(self._val(report, "output_torque"), "349")
        torque_field = next(f for f in confirmed["fields"] if f["key"] == "output_torque")
        self.assertEqual(torque_field["original_value"], "349")
        self.assertEqual(torque_field["status"], "Edited")

        pdf_path = self.scratch / "test_from_excel.pdf"
        generate_pdf(confirmed, pdf_path)
        self.assertTrue(pdf_path.exists())
        if fitz is not None:
            text = fitz.open(pdf_path)[0].get_text()
            self.assertIn("337.5", text)
            self.assertNotIn("(Nm) : 349", text)
            self.assertNotIn("Original_Data", text)
            self.assertNotIn("Change_Log", text)
            self.assertIn("DS-021-0046", text)
            self.assertIn("Product Information", text)

    def test_blank_required_value_fails_validation(self):
        report = self._report()
        path = self.scratch / "test_blank.xlsx"
        generate_excel(report, path)
        wb = load_workbook(path)
        ws = wb[SHEET_EDITABLE]
        keys = []
        row = DATA_START_ROW
        while ws.cell(row, 1).value:
            keys.append(ws.cell(row, 1).value)
            row += 1
        r = keys.index("output_torque") + DATA_START_ROW
        ws.cell(r, 3).value = ""
        wb.save(path)
        loaded = load_excel(path)
        check = validate_rows(
            loaded["editable"],
            loaded["original"],
            [f["key"] for f in report["fields"]],
        )
        self.assertFalse(check["ok"])
        self.assertTrue(any("Output torque" in e for e in check["errors"]))

    def test_new_product_fingerprint_differs(self):
        a = self.engine.lookup("RF", 1.5, 20.0, model="RF 042 - IEC 90 B5")["report"]
        b = self.engine.lookup("RF", 1.5, 36.53)["report"]
        self.assertNotEqual(selection_fingerprint(a), selection_fingerprint(b))
        self.assertNotEqual(self._val(a, "output_torque"), self._val(b, "output_torque"))


if __name__ == "__main__":
    unittest.main()
